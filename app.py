#!/usr/bin/env python3
"""
Siyafundza LMS – Flask REST API + React SPA frontend
"""
import os, sqlite3, json, re, io, secrets
from docx import Document
from pypdf import PdfReader
from datetime import datetime
from functools import wraps
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from flask import (
    Flask, request, jsonify, send_from_directory, g, session, send_file
)
from flask_login import (
    LoginManager, UserMixin, login_user, logout_user,
    login_required, current_user
)
from flask_cors import CORS

app = Flask(__name__, static_folder="static", static_url_path="/static")
_base_dir = os.path.dirname(os.path.abspath(__file__))
IS_PRODUCTION = bool(os.environ.get("RENDER")) or os.environ.get("FLASK_ENV") == "production"

_secret = os.environ.get("SECRET_KEY")
if not _secret:
    if IS_PRODUCTION:
        raise RuntimeError("SECRET_KEY environment variable must be set in production")
    _secret = "dev-only-insecure-key"
app.config["SECRET_KEY"] = _secret
# Uploads live outside /static so files are only reachable through the authenticated API
app.config["UPLOAD_FOLDER"] = os.environ.get("UPLOAD_FOLDER", os.path.join(_base_dir, "instance", "uploads"))
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
# Portable DB path: project/instance/lms.db (override with LMS_DB env var)
_instance_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "instance")
os.makedirs(_instance_dir, exist_ok=True)
app.config["DATABASE"] = os.environ.get("LMS_DB", os.path.join(_instance_dir, "lms.db"))
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SECURE"] = IS_PRODUCTION
app.config["REMEMBER_COOKIE_SECURE"] = IS_PRODUCTION

# The SPA is served by this same app, so CORS is only needed if you host the frontend elsewhere.
_cors_origins = [o.strip() for o in os.environ.get("CORS_ORIGINS", "").split(",") if o.strip()]
if _cors_origins:
    CORS(app, supports_credentials=True, origins=_cors_origins)

ALLOWED_EXTENSIONS = {"pdf", "doc", "docx", "txt", "png", "jpg", "jpeg", "ppt", "pptx", "xlsx"}

login_manager = LoginManager()
login_manager.init_app(app)


def get_db():
    if "db" not in g:
        db_path = app.config["DATABASE"]
        parent = os.path.dirname(db_path)
        if parent:
            os.makedirs(parent, exist_ok=True)
        g.db = sqlite3.connect(db_path)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def row_to_dict(row):
    if row is None:
        return None
    return dict(row)


def init_db():
    db = get_db()
    db.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('student', 'teacher')),
            is_admin INTEGER NOT NULL DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS resources (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            filename TEXT,
            original_filename TEXT,
            resource_type TEXT NOT NULL DEFAULT 'notes',
            uploaded_by INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (uploaded_by) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS tests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            total_marks INTEGER NOT NULL DEFAULT 50,
            duration_minutes INTEGER DEFAULT 60,
            is_published INTEGER DEFAULT 0,
            created_by INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            test_id INTEGER NOT NULL,
            question_text TEXT NOT NULL,
            question_type TEXT NOT NULL CHECK(question_type IN ('mcq', 'short', 'essay')),
            options TEXT,
            correct_answer TEXT,
            marks INTEGER NOT NULL DEFAULT 1,
            order_num INTEGER DEFAULT 0,
            FOREIGN KEY (test_id) REFERENCES tests(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            test_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL,
            score REAL,
            max_score INTEGER,
            status TEXT DEFAULT 'in_progress',
            submitted_at TIMESTAMP,
            graded_at TIMESTAMP,
            graded_by INTEGER,
            feedback TEXT,
            FOREIGN KEY (test_id) REFERENCES tests(id),
            FOREIGN KEY (student_id) REFERENCES users(id),
            FOREIGN KEY (graded_by) REFERENCES users(id),
            UNIQUE(test_id, student_id)
        );
        CREATE TABLE IF NOT EXISTS answers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            submission_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            answer_text TEXT,
            marks_awarded REAL,
            FOREIGN KEY (submission_id) REFERENCES submissions(id) ON DELETE CASCADE,
            FOREIGN KEY (question_id) REFERENCES questions(id)
        );
    """)
    db.commit()

    # Migration for databases created before the is_admin column existed
    cols = [r["name"] for r in db.execute("PRAGMA table_info(users)").fetchall()]
    if "is_admin" not in cols:
        db.execute("ALTER TABLE users ADD COLUMN is_admin INTEGER NOT NULL DEFAULT 0")
        db.commit()

    # Default teacher/admin: the only account that can create other teachers.
    admin_username = os.environ.get("ADMIN_USERNAME", "admin").strip() or "admin"
    admin_password = os.environ.get("ADMIN_PASSWORD", "")
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@siyafundza.school").strip()
    admin_name = os.environ.get("ADMIN_FULL_NAME", "Administrator").strip()

    has_admin = db.execute("SELECT COUNT(*) AS c FROM users WHERE is_admin = 1").fetchone()["c"]
    if not has_admin:
        existing = db.execute("SELECT id FROM users WHERE username = ?", (admin_username,)).fetchone()
        if existing:
            db.execute("UPDATE users SET is_admin = 1, role = 'teacher' WHERE id = ?", (existing["id"],))
        else:
            if not admin_password:
                if IS_PRODUCTION:
                    raise RuntimeError("ADMIN_PASSWORD environment variable must be set to create the default admin")
                admin_password = secrets.token_urlsafe(12)
                print(f"[dev] Generated admin password for '{admin_username}': {admin_password}")
            db.execute(
                "INSERT INTO users (username, email, password_hash, full_name, role, is_admin) VALUES (?,?,?,?,?,1)",
                (admin_username, admin_email, generate_password_hash(admin_password), admin_name, "teacher"),
            )
        db.commit()
        print(f"Default admin ready: '{admin_username}'")


class User(UserMixin):
    def __init__(self, id, username, email, full_name, role, is_admin=False):
        self.id = id
        self.is_admin = bool(is_admin)
        self.username = username
        self.email = email
        self.full_name = full_name
        self.role = role

    def to_dict(self):
        return {
            "id": self.id,
            "username": self.username,
            "email": self.email,
            "full_name": self.full_name,
            "role": self.role,
            "is_admin": self.is_admin,
        }


@login_manager.user_loader
def load_user(user_id):
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if row:
        return User(row["id"], row["username"], row["email"], row["full_name"], row["role"], row["is_admin"])
    return None


@login_manager.unauthorized_handler
def unauthorized():
    return jsonify({"error": "Authentication required"}), 401


def teacher_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != "teacher":
            return jsonify({"error": "Teacher access required"}), 403
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            return jsonify({"error": "Administrator access required"}), 403
        return f(*args, **kwargs)
    return decorated


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


# ── SPA entry ──────────────────────────────────────────────────────────────
@app.route("/")
@app.route("/login")
@app.route("/register")
@app.route("/resources")
@app.route("/tests")
@app.route("/tests/create")
@app.route("/tests/<path:sub>")
@app.route("/submissions")
@app.route("/submissions/<path:sub>")
@app.route("/results/<path:sub>")
@app.route("/users")
def spa(sub=None):
    return send_from_directory(app.static_folder, "index.html")


# ── Auth API ───────────────────────────────────────────────────────────────
@app.route("/api/me")
def api_me():
    if current_user.is_authenticated:
        return jsonify({"user": current_user.to_dict()})
    return jsonify({"user": None})


@app.route("/api/login", methods=["POST"])
def api_login():
    data = request.get_json() or {}
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    db = get_db()
    row = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    if row and check_password_hash(row["password_hash"], password):
        user = User(row["id"], row["username"], row["email"], row["full_name"], row["role"], row["is_admin"])
        login_user(user)
        return jsonify({"user": user.to_dict(), "message": f"Welcome, {user.full_name}!"})
    return jsonify({"error": "Invalid username or password"}), 401


@app.route("/api/register", methods=["POST"])
def api_register():
    data = request.get_json() or {}
    username = (data.get("username") or "").strip()
    email = (data.get("email") or "").strip()
    full_name = (data.get("full_name") or "").strip()
    password = data.get("password") or ""
    # Public registration always creates students. Teachers are created by the admin only.
    role = "student"
    if not all([username, email, full_name, password]):
        return jsonify({"error": "All fields are required"}), 400
    if len(password) < 6:
        return jsonify({"error": "Password must be at least 6 characters"}), 400
    db = get_db()
    try:
        db.execute(
            "INSERT INTO users (username, email, password_hash, full_name, role) VALUES (?,?,?,?,?)",
            (username, email, generate_password_hash(password), full_name, role),
        )
        db.commit()
        return jsonify({"message": "Registration successful"})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Username or email already exists"}), 400


@app.route("/api/logout", methods=["POST"])
@login_required
def api_logout():
    logout_user()
    return jsonify({"message": "Logged out"})


# ── Dashboard ──────────────────────────────────────────────────────────────
@app.route("/api/dashboard")
@login_required
def api_dashboard():
    db = get_db()
    resources = [row_to_dict(r) for r in db.execute(
        """SELECT r.*, u.full_name AS uploader_name
           FROM resources r JOIN users u ON r.uploaded_by = u.id
           ORDER BY r.created_at DESC LIMIT 5"""
    ).fetchall()]
    if current_user.role == "teacher":
        tests = [row_to_dict(r) for r in db.execute(
            "SELECT * FROM tests WHERE created_by = ? ORDER BY created_at DESC LIMIT 5",
            (current_user.id,),
        ).fetchall()]
        pending = [row_to_dict(r) for r in db.execute(
            """SELECT s.*, t.title AS test_title, u.full_name AS student_name
               FROM submissions s
               JOIN tests t ON s.test_id = t.id
               JOIN users u ON s.student_id = u.id
               WHERE t.created_by = ? AND s.status = 'submitted'
               ORDER BY s.submitted_at DESC LIMIT 10""",
            (current_user.id,),
        ).fetchall()]
        return jsonify({"resources": resources, "tests": tests, "pending": pending})
    else:
        tests = [row_to_dict(r) for r in db.execute(
            "SELECT * FROM tests WHERE is_published = 1 ORDER BY created_at DESC LIMIT 5"
        ).fetchall()]
        my_subs = [row_to_dict(r) for r in db.execute(
            """SELECT s.*, t.title AS test_title
               FROM submissions s JOIN tests t ON s.test_id = t.id
               WHERE s.student_id = ? ORDER BY s.submitted_at DESC LIMIT 5""",
            (current_user.id,),
        ).fetchall()]
        return jsonify({"resources": resources, "tests": tests, "my_subs": my_subs})


# ── Resources ──────────────────────────────────────────────────────────────
@app.route("/api/resources")
@login_required
def api_resources():
    db = get_db()
    rtype = request.args.get("type", "")
    if rtype:
        rows = db.execute(
            """SELECT r.*, u.full_name AS uploader_name
               FROM resources r JOIN users u ON r.uploaded_by = u.id
               WHERE r.resource_type = ? ORDER BY r.created_at DESC""",
            (rtype,),
        ).fetchall()
    else:
        rows = db.execute(
            """SELECT r.*, u.full_name AS uploader_name
               FROM resources r JOIN users u ON r.uploaded_by = u.id
               ORDER BY r.created_at DESC"""
        ).fetchall()
    return jsonify({"resources": [row_to_dict(r) for r in rows]})


@app.route("/api/resources", methods=["POST"])
@login_required
def api_upload_resource():
    title = (request.form.get("title") or "").strip()
    description = (request.form.get("description") or "").strip()
    resource_type = request.form.get("resource_type", "notes")
    if not title:
        return jsonify({"error": "Title is required"}), 400
    file = request.files.get("file")
    filename = original = None
    if file and file.filename and allowed_file(file.filename):
        original = secure_filename(file.filename)
        filename = f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{original}"
        folder = os.path.join(app.config["UPLOAD_FOLDER"], "resources")
        os.makedirs(folder, exist_ok=True)
        file.save(os.path.join(folder, filename))
    db = get_db()
    db.execute(
        """INSERT INTO resources (title, description, filename, original_filename, resource_type, uploaded_by)
           VALUES (?,?,?,?,?,?)""",
        (title, description, filename, original, resource_type, current_user.id),
    )
    db.commit()
    return jsonify({"message": "Resource shared successfully"})


@app.route("/api/resources/<int:rid>/download")
@login_required
def api_download_resource(rid):
    db = get_db()
    row = db.execute("SELECT * FROM resources WHERE id = ?", (rid,)).fetchone()
    if not row or not row["filename"]:
        return jsonify({"error": "Not found"}), 404
    return send_from_directory(
        os.path.join(app.config["UPLOAD_FOLDER"], "resources"),
        row["filename"],
        as_attachment=True,
        download_name=row["original_filename"] or row["filename"],
    )


@app.route("/api/resources/<int:rid>", methods=["DELETE"])
@login_required
def api_delete_resource(rid):
    db = get_db()
    row = db.execute("SELECT * FROM resources WHERE id = ?", (rid,)).fetchone()
    if not row:
        return jsonify({"error": "Not found"}), 404
    if row["uploaded_by"] != current_user.id and current_user.role != "teacher":
        return jsonify({"error": "Forbidden"}), 403
    if row["filename"]:
        path = os.path.join(app.config["UPLOAD_FOLDER"], "resources", row["filename"])
        if os.path.exists(path):
            os.remove(path)
    db.execute("DELETE FROM resources WHERE id = ?", (rid,))
    db.commit()
    return jsonify({"message": "Deleted"})


# ── Tests ──────────────────────────────────────────────────────────────────
@app.route("/api/tests")
@login_required
def api_tests():
    db = get_db()
    if current_user.role == "teacher":
        rows = db.execute(
            "SELECT * FROM tests WHERE created_by = ? ORDER BY created_at DESC",
            (current_user.id,),
        ).fetchall()
    else:
        rows = db.execute(
            "SELECT * FROM tests WHERE is_published = 1 ORDER BY created_at DESC"
        ).fetchall()
    return jsonify({"tests": [row_to_dict(r) for r in rows]})


@app.route("/api/tests", methods=["POST"])
@login_required
@teacher_required
def api_create_test():
    data = request.get_json() or {}
    title = (data.get("title") or "").strip()
    description = (data.get("description") or "").strip()
    total_marks = min(int(data.get("total_marks", 50)), 50)
    duration = int(data.get("duration_minutes", 60))
    if not title:
        return jsonify({"error": "Title is required"}), 400
    db = get_db()
    cur = db.execute(
        """INSERT INTO tests (title, description, total_marks, duration_minutes, created_by)
           VALUES (?,?,?,?,?)""",
        (title, description, total_marks, duration, current_user.id),
    )
    db.commit()
    test = row_to_dict(db.execute("SELECT * FROM tests WHERE id = ?", (cur.lastrowid,)).fetchone())
    return jsonify({"test": test, "message": "Test created"})






def extract_text_from_pdf(file_storage):
    reader = PdfReader(file_storage)
    parts = []
    for page in reader.pages:
        t = page.extract_text() or ""
        parts.append(t)
    return "\n".join(parts)


def extract_text_from_docx(file_storage):
    data = file_storage.read()
    doc = Document(io.BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    # also tables
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)



def parse_questions_from_text(text, default_marks=5):
    """
    Heuristic parser for exam-style text.
    Skips instruction/header blocks; only keeps numbered questions.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n\s*([A-Da-d])[\).\:\-]\s+", r"\n\1) ", text)

    # Lines that look like instructions / headers (not questions)
    instr_re = re.compile(
        r"(?i)^("
        r"instructions?|directions?|read\s+the\s+following|answer\s+all|"
        r"total\s+marks?|time\s*allowed|duration|name\s*:|date\s*:|class\s*:|"
        r"grade\s*:|subject\s*:|school\s*:|section\s*:|student\s+name|"
        r"do\s+not\s+open|write\s+your\s+name|use\s+blue|fill\s+in|"
        r"choose\s+the\s+correct|circle\s+the\s+correct|match\s+the\s+following|"
        r"section\s+[a-z]|part\s+[a-z]|paper\s+\d+|term\s+\w+"
        r").*"
    )

    pattern = re.compile(
        r"(?:^|\n)\s*(?:Q(?:uestion)?\s*)?(\d+)\s*[\.\)\-]\s+",
        re.IGNORECASE,
    )
    indices = [(m.start(), m.end(), m.group(1)) for m in pattern.finditer(text)]
    questions = []

    if not indices:
        return questions  # do not invent an essay from instructions

    for i, (start_i, end_i, num) in enumerate(indices):
        block_end = indices[i + 1][0] if i + 1 < len(indices) else len(text)
        block = text[end_i:block_end].strip()
        if not block:
            continue

        # Skip blocks that are mostly instructions
        first_line = block.split("\n")[0].strip()
        if instr_re.match(first_line) and len(block) < 200:
            continue
        # Skip very short non-question fragments
        if len(re.sub(r"\s+", " ", block).strip()) < 8:
            continue

        marks = default_marks
        m = re.search(r"[\(\[]\s*(\d+)\s*(?:marks?|pts?|points?)?\s*[\)\]]", block, re.I)
        if m:
            marks = max(1, int(m.group(1)))
            block = (block[: m.start()] + block[m.end() :]).strip()

        correct = ""
        am = re.search(r"(?im)^\s*(?:correct\s*)?answer\s*[:=\-]\s*(.+?)\s*$", block)
        if am:
            correct = am.group(1).strip()
            if re.fullmatch(r"[A-Da-d]", correct):
                correct = correct.upper()
            block = (block[: am.start()] + block[am.end() :]).strip()

        opt_pattern = re.compile(r"(?m)^\s*([A-D])\)\s*(.+)$")
        opts_found = opt_pattern.findall(block)
        options = []
        if opts_found:
            qtext = opt_pattern.sub("", block).strip()
            qtext = re.sub(r"\n{2,}", "\n", qtext).strip()
            opt_map = {k.upper(): v.strip() for k, v in opts_found}
            options = [opt_map.get(L, "") for L in ("A", "B", "C", "D")]
            qtype = "mcq"
            if correct and len(correct) == 1:
                correct = correct.upper()
        else:
            qtext = block.strip()
            if len(qtext) < 120 and not re.search(
                r"\b(explain|describe|discuss|essay|paragraph|write\s+about)\b", qtext, re.I
            ):
                qtype = "short"
            else:
                qtype = "essay"

        # Strip leftover instruction phrases from question text
        qtext = re.sub(r"(?i)^(instructions?|directions?)\s*[:.\-]\s*", "", qtext).strip()
        qtext = re.sub(r"\s+", " ", qtext).strip()
        if not qtext or instr_re.match(qtext):
            continue
        # Avoid treating "Answer all questions" style lines as questions
        if re.match(r"(?i)^(answer|choose|select|write|fill|match)\b", qtext) and len(qtext) < 60:
            continue

        questions.append(
            {
                "question_text": qtext[:2000],
                "question_type": qtype,
                "marks": marks,
                "correct_answer": correct,
                "options": options,
            }
        )

    return questions


@app.route("/api/tests/parse-file", methods=["POST"])
@login_required
@teacher_required
def api_parse_test_file():
    """Upload PDF/DOCX/JSON/TXT and return parsed test JSON for preview."""
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file uploaded"}), 400

    filename = f.filename.lower()
    total_marks = min(int(request.form.get("total_marks") or 50), 50)
    duration = int(request.form.get("duration_minutes") or 60)
    title = (request.form.get("title") or "").strip()
    default_marks = max(1, int(request.form.get("default_marks") or 5))

    try:
        if filename.endswith(".json"):
            data = json.loads(f.read().decode("utf-8"))
            return jsonify({"test": data, "source": "json", "raw_text": ""})

        if filename.endswith(".pdf"):
            raw = extract_text_from_pdf(f.stream)
        elif filename.endswith(".docx"):
            raw = extract_text_from_docx(f)
        elif filename.endswith(".txt") or filename.endswith(".md"):
            raw = f.read().decode("utf-8", errors="replace")
        else:
            return jsonify({"error": "Unsupported file type. Use PDF, DOCX, JSON, or TXT."}), 400

        raw = (raw or "").strip()
        if not raw:
            return jsonify({"error": "No text could be extracted from this file. If it is a scanned PDF, OCR is required."}), 400

        questions = parse_questions_from_text(raw, default_marks=default_marks)
        if not title:
            # First non-empty line as title guess
            for line in raw.split("\n"):
                line = line.strip()
                if len(line) > 5:
                    title = line[:120]
                    break
            if not title:
                title = "Imported Test"

        # Cap marks to total_marks
        used = 0
        for q in questions:
            if used + q["marks"] > total_marks:
                q["marks"] = max(1, total_marks - used)
            used += q["marks"]
            if used >= total_marks:
                # drop remaining
                idx = questions.index(q)
                questions = questions[: idx + 1]
                break

        test = {
            "title": title,
            "description": f"Imported from {f.filename}",
            "total_marks": total_marks,
            "duration_minutes": duration,
            "questions": questions,
        }
        return jsonify({
            "test": test,
            "source": filename.rsplit(".", 1)[-1],
            "raw_text": raw[:5000],
            "questions_found": len(questions),
        })
    except Exception as e:
        return jsonify({"error": "Failed to parse file: " + str(e)}), 400


@app.route("/api/tests/import", methods=["POST"])
@login_required
@teacher_required
def api_import_test():
    """Import a test from JSON body or uploaded .json file.
    Expected JSON shape:
    {
      "title": "...",
      "description": "...",
      "total_marks": 50,
      "duration_minutes": 60,
      "questions": [
        {
          "question_text": "...",
          "question_type": "mcq|short|essay",
          "marks": 5,
          "correct_answer": "A",
          "options": ["optA", "optB", "optC", "optD"]
        }
      ]
    }
    """
    data = None
    if request.files.get("file"):
        f = request.files["file"]
        try:
            data = json.loads(f.read().decode("utf-8"))
        except Exception as e:
            return jsonify({"error": "Invalid JSON file: " + str(e)}), 400
    else:
        data = request.get_json() or {}

    title = (data.get("title") or "").strip()
    if not title:
        return jsonify({"error": "Title is required in the import file"}), 400

    description = (data.get("description") or "").strip()
    total_marks = min(int(data.get("total_marks") or 50), 50)
    duration = int(data.get("duration_minutes") or 60)
    questions = data.get("questions") or []

    if not questions:
        return jsonify({"error": "Import must include at least one question"}), 400

    db = get_db()
    cur = db.execute(
        """INSERT INTO tests (title, description, total_marks, duration_minutes, created_by)
           VALUES (?,?,?,?,?)""",
        (title, description, total_marks, duration, current_user.id),
    )
    test_id = cur.lastrowid

    used = 0
    order = 0
    for q in questions:
        qtext = (q.get("question_text") or "").strip()
        if not qtext:
            continue
        qtype = q.get("question_type") or "mcq"
        if qtype not in ("mcq", "short", "essay"):
            qtype = "short"
        marks = max(1, int(q.get("marks") or 1))
        if used + marks > total_marks:
            marks = max(1, total_marks - used)
        correct = (q.get("correct_answer") or "").strip()
        options = q.get("options")
        opt_json = None
        if qtype == "mcq":
            if isinstance(options, list):
                # pad to 4
                opts = (options + ["", "", "", ""])[:4]
                opt_json = json.dumps(opts)
            else:
                opt_json = json.dumps([
                    q.get("opt_a") or "",
                    q.get("opt_b") or "",
                    q.get("opt_c") or "",
                    q.get("opt_d") or "",
                ])
        order += 1
        used += marks
        db.execute(
            """INSERT INTO questions (test_id, question_text, question_type, options, correct_answer, marks, order_num)
               VALUES (?,?,?,?,?,?,?)""",
            (test_id, qtext, qtype, opt_json, correct, marks, order),
        )

    db.commit()
    test = row_to_dict(db.execute("SELECT * FROM tests WHERE id = ?", (test_id,)).fetchone())
    return jsonify({"test": test, "message": f"Imported test with {order} questions", "questions_imported": order})


@app.route("/api/tests/<int:test_id>")
@login_required
@teacher_required
def api_get_test(test_id):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    questions = db.execute(
        "SELECT * FROM questions WHERE test_id = ? ORDER BY order_num, id", (test_id,)
    ).fetchall()
    return jsonify({"test": row_to_dict(test), "questions": [row_to_dict(q) for q in questions]})


@app.route("/api/tests/<int:test_id>/questions", methods=["POST"])
@login_required
@teacher_required
def api_add_question(test_id):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    data = request.get_json() or {}
    qtext = (data.get("question_text") or "").strip()
    qtype = data.get("question_type", "mcq")
    marks = int(data.get("marks", 1))
    correct = (data.get("correct_answer") or "").strip()
    if not qtext or marks < 1:
        return jsonify({"error": "Invalid question"}), 400
    existing = db.execute(
        "SELECT COALESCE(SUM(marks),0) AS s FROM questions WHERE test_id = ?", (test_id,)
    ).fetchone()["s"]
    if existing + marks > test["total_marks"]:
        return jsonify({"error": f"Cannot exceed total marks of {test['total_marks']}. Current: {existing}"}), 400
    options = None
    if qtype == "mcq":
        opts = [data.get("opt_a", ""), data.get("opt_b", ""), data.get("opt_c", ""), data.get("opt_d", "")]
        options = json.dumps(opts)
    order = db.execute("SELECT COUNT(*) AS c FROM questions WHERE test_id = ?", (test_id,)).fetchone()["c"] + 1
    db.execute(
        """INSERT INTO questions (test_id, question_text, question_type, options, correct_answer, marks, order_num)
           VALUES (?,?,?,?,?,?,?)""",
        (test_id, qtext, qtype, options, correct, marks, order),
    )
    db.commit()
    return jsonify({"message": "Question added"})




@app.route("/api/tests/<int:test_id>/questions/<int:qid>", methods=["PUT"])
@login_required
@teacher_required
def api_update_question(test_id, qid):
    """Update question text, type, marks, options, or correct_answer (for auto-grading)."""
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    q = db.execute(
        "SELECT * FROM questions WHERE id = ? AND test_id = ?", (qid, test_id)
    ).fetchone()
    if not q:
        return jsonify({"error": "Question not found"}), 404

    data = request.get_json() or {}
    qtext = (data.get("question_text") if data.get("question_text") is not None else q["question_text"]).strip()
    qtype = data.get("question_type") or q["question_type"]
    if qtype not in ("mcq", "short", "essay"):
        qtype = q["question_type"]
    marks = int(data.get("marks") if data.get("marks") is not None else q["marks"])
    correct = data.get("correct_answer") if data.get("correct_answer") is not None else (q["correct_answer"] or "")
    correct = (correct or "").strip()

    options = q["options"]
    if qtype == "mcq":
        if "options" in data and isinstance(data["options"], list):
            opts = (data["options"] + ["", "", "", ""])[:4]
            options = json.dumps(opts)
        elif any(k in data for k in ("opt_a", "opt_b", "opt_c", "opt_d")):
            options = json.dumps([
                data.get("opt_a", ""),
                data.get("opt_b", ""),
                data.get("opt_c", ""),
                data.get("opt_d", ""),
            ])
        if correct and len(correct) == 1:
            correct = correct.upper()

    # marks budget check (excluding this question's old marks)
    others = db.execute(
        "SELECT COALESCE(SUM(marks),0) AS s FROM questions WHERE test_id = ? AND id != ?",
        (test_id, qid),
    ).fetchone()["s"]
    if others + marks > test["total_marks"]:
        return jsonify({"error": f"Total would exceed {test['total_marks']} marks"}), 400

    db.execute(
        """UPDATE questions SET question_text=?, question_type=?, options=?, correct_answer=?, marks=?
           WHERE id=? AND test_id=?""",
        (qtext, qtype, options, correct, marks, qid, test_id),
    )
    db.commit()
    return jsonify({"message": "Question updated"})


@app.route("/api/tests/<int:test_id>/questions/<int:qid>", methods=["DELETE"])
@login_required
@teacher_required
def api_delete_question(test_id, qid):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    db.execute("DELETE FROM questions WHERE id = ? AND test_id = ?", (qid, test_id))
    db.commit()
    return jsonify({"message": "Deleted"})


@app.route("/api/tests/<int:test_id>/publish", methods=["POST"])
@login_required
@teacher_required
def api_publish(test_id):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    count = db.execute("SELECT COUNT(*) AS c FROM questions WHERE test_id = ?", (test_id,)).fetchone()["c"]
    if count == 0:
        return jsonify({"error": "Add at least one question before publishing"}), 400
    db.execute("UPDATE tests SET is_published = 1 WHERE id = ?", (test_id,))
    db.commit()
    return jsonify({"message": "Published"})


@app.route("/api/tests/<int:test_id>/unpublish", methods=["POST"])
@login_required
@teacher_required
def api_unpublish(test_id):
    db = get_db()
    db.execute(
        "UPDATE tests SET is_published = 0 WHERE id = ? AND created_by = ?",
        (test_id, current_user.id),
    )
    db.commit()
    return jsonify({"message": "Unpublished"})


@app.route("/api/tests/<int:test_id>", methods=["DELETE"])
@login_required
@teacher_required
def api_delete_test(test_id):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND created_by = ?", (test_id, current_user.id)
    ).fetchone()
    if not test:
        return jsonify({"error": "Not found"}), 404
    db.execute("DELETE FROM answers WHERE submission_id IN (SELECT id FROM submissions WHERE test_id = ?)", (test_id,))
    db.execute("DELETE FROM submissions WHERE test_id = ?", (test_id,))
    db.execute("DELETE FROM questions WHERE test_id = ?", (test_id,))
    db.execute("DELETE FROM tests WHERE id = ?", (test_id,))
    db.commit()
    return jsonify({"message": "Deleted"})


# ── Take & submit test ─────────────────────────────────────────────────────


@app.route("/api/tests/<int:test_id>/draft", methods=["POST"])
@login_required
def api_save_draft(test_id):
    """Auto-save in-progress answers (does not finalise submission)."""
    db = get_db()
    test = db.execute("SELECT * FROM tests WHERE id = ? AND is_published = 1", (test_id,)).fetchone()
    if not test:
        return jsonify({"error": "Test not available"}), 404
    existing = db.execute(
        "SELECT * FROM submissions WHERE test_id = ? AND student_id = ?",
        (test_id, current_user.id),
    ).fetchone()
    if existing and existing["status"] in ("submitted", "graded"):
        return jsonify({"error": "Already submitted"}), 400

    data = request.get_json() or {}
    answers_map = data.get("answers") or {}
    answers_map = {str(k): v for k, v in answers_map.items()}

    if existing:
        sub_id = existing["id"]
        db.execute(
            "UPDATE submissions SET status = 'in_progress' WHERE id = ?",
            (sub_id,),
        )
        db.execute("DELETE FROM answers WHERE submission_id = ?", (sub_id,))
    else:
        cur = db.execute(
            """INSERT INTO submissions (test_id, student_id, max_score, status)
               VALUES (?,?,?,'in_progress')""",
            (test_id, current_user.id, test["total_marks"]),
        )
        sub_id = cur.lastrowid

    questions = db.execute(
        "SELECT id FROM questions WHERE test_id = ?", (test_id,)
    ).fetchall()
    for q in questions:
        ans = (answers_map.get(str(q["id"])) or "").strip()
        if ans:
            db.execute(
                "INSERT INTO answers (submission_id, question_id, answer_text, marks_awarded) VALUES (?,?,?,NULL)",
                (sub_id, q["id"], ans),
            )
    db.commit()
    return jsonify({"message": "Draft saved", "submission_id": sub_id})


@app.route("/api/tests/<int:test_id>/draft", methods=["GET"])
@login_required
def api_get_draft(test_id):
    """Load previously auto-saved answers for a test."""
    db = get_db()
    existing = db.execute(
        "SELECT * FROM submissions WHERE test_id = ? AND student_id = ?",
        (test_id, current_user.id),
    ).fetchone()
    if not existing:
        return jsonify({"answers": {}, "status": None})
    if existing["status"] in ("submitted", "graded"):
        return jsonify({"error": "Already submitted", "submission_id": existing["id"], "status": existing["status"]}), 400
    rows = db.execute(
        "SELECT question_id, answer_text FROM answers WHERE submission_id = ?",
        (existing["id"],),
    ).fetchall()
    answers = {str(r["question_id"]): r["answer_text"] or "" for r in rows}
    return jsonify({"answers": answers, "status": existing["status"], "submission_id": existing["id"]})


@app.route("/api/tests/<int:test_id>/take")
@login_required
def api_take_test(test_id):
    db = get_db()
    test = db.execute("SELECT * FROM tests WHERE id = ? AND is_published = 1", (test_id,)).fetchone()
    if not test:
        return jsonify({"error": "Test not available"}), 404
    existing = db.execute(
        "SELECT * FROM submissions WHERE test_id = ? AND student_id = ?",
        (test_id, current_user.id),
    ).fetchone()
    if existing and existing["status"] in ("submitted", "graded"):
        return jsonify({"error": "You have already submitted this test", "submission_id": existing["id"]}), 400
    questions = db.execute(
        "SELECT id, question_text, question_type, options, marks, order_num FROM questions WHERE test_id = ? ORDER BY order_num, id",
        (test_id,),
    ).fetchall()
    return jsonify({"test": row_to_dict(test), "questions": [row_to_dict(q) for q in questions]})


@app.route("/api/tests/<int:test_id>/submit", methods=["POST"])
@login_required
def api_submit_test(test_id):
    db = get_db()
    test = db.execute("SELECT * FROM tests WHERE id = ? AND is_published = 1", (test_id,)).fetchone()
    if not test:
        return jsonify({"error": "Test not available"}), 404
    existing = db.execute(
        "SELECT * FROM submissions WHERE test_id = ? AND student_id = ?",
        (test_id, current_user.id),
    ).fetchone()
    if existing and existing["status"] in ("submitted", "graded"):
        return jsonify({"error": "Already submitted", "submission_id": existing["id"]}), 400

    data = request.get_json() or {}
    answers_map = data.get("answers") or {}
    # keys may be strings from JSON
    answers_map = {str(k): v for k, v in answers_map.items()}

    questions = db.execute(
        "SELECT * FROM questions WHERE test_id = ? ORDER BY order_num, id", (test_id,)
    ).fetchall()

    if existing:
        sub_id = existing["id"]
        db.execute(
            "UPDATE submissions SET status='submitted', submitted_at=? WHERE id=?",
            (datetime.utcnow().isoformat(), sub_id),
        )
        db.execute("DELETE FROM answers WHERE submission_id = ?", (sub_id,))
    else:
        cur = db.execute(
            """INSERT INTO submissions (test_id, student_id, max_score, status, submitted_at)
               VALUES (?,?,?,'submitted',?)""",
            (test_id, current_user.id, test["total_marks"], datetime.utcnow().isoformat()),
        )
        sub_id = cur.lastrowid

    auto_score = 0.0
    needs_manual = False
    for q in questions:
        ans = (answers_map.get(str(q["id"])) or "").strip()
        marks_awarded = None
        if q["question_type"] == "mcq":
            marks_awarded = q["marks"] if ans.lower() == (q["correct_answer"] or "").lower() else 0
            auto_score += marks_awarded
        elif q["question_type"] == "short":
            marks_awarded = q["marks"] if ans.lower() == (q["correct_answer"] or "").lower() else 0
            auto_score += marks_awarded
        else:
            needs_manual = True
        db.execute(
            "INSERT INTO answers (submission_id, question_id, answer_text, marks_awarded) VALUES (?,?,?,?)",
            (sub_id, q["id"], ans, marks_awarded),
        )

    status = "submitted" if needs_manual else "graded"
    final_score = None if needs_manual else auto_score
    db.execute(
        "UPDATE submissions SET score=?, status=?, graded_at=? WHERE id=?",
        (final_score, status, datetime.utcnow().isoformat() if status == "graded" else None, sub_id),
    )
    db.commit()
    return jsonify({"submission_id": sub_id, "message": "Submitted", "auto_score": auto_score if not needs_manual else None})


# ── Submissions & grading ──────────────────────────────────────────────────
@app.route("/api/submissions")
@login_required
@teacher_required
def api_submissions():
    db = get_db()
    rows = db.execute(
        """SELECT s.*, t.title AS test_title, u.full_name AS student_name
           FROM submissions s
           JOIN tests t ON s.test_id = t.id
           JOIN users u ON s.student_id = u.id
           WHERE t.created_by = ?
           ORDER BY s.submitted_at DESC""",
        (current_user.id,),
    ).fetchall()
    return jsonify({"submissions": [row_to_dict(r) for r in rows]})


@app.route("/api/submissions/<int:sid>")
@login_required
@teacher_required
def api_get_submission(sid):
    db = get_db()
    sub = db.execute(
        """SELECT s.*, t.title AS test_title, t.total_marks, t.created_by, u.full_name AS student_name
           FROM submissions s
           JOIN tests t ON s.test_id = t.id
           JOIN users u ON s.student_id = u.id
           WHERE s.id = ?""",
        (sid,),
    ).fetchone()
    if not sub or sub["created_by"] != current_user.id:
        return jsonify({"error": "Not found"}), 404
    answers = db.execute(
        """SELECT a.*, q.question_text, q.question_type, q.marks AS max_marks, q.correct_answer, q.options
           FROM answers a JOIN questions q ON a.question_id = q.id
           WHERE a.submission_id = ?
           ORDER BY q.order_num, q.id""",
        (sid,),
    ).fetchall()
    return jsonify({"submission": row_to_dict(sub), "answers": [row_to_dict(a) for a in answers]})


@app.route("/api/submissions/<int:sid>/grade", methods=["POST"])
@login_required
@teacher_required
def api_grade(sid):
    db = get_db()
    sub = db.execute(
        """SELECT s.*, t.total_marks, t.created_by
           FROM submissions s JOIN tests t ON s.test_id = t.id WHERE s.id = ?""",
        (sid,),
    ).fetchone()
    if not sub or sub["created_by"] != current_user.id:
        return jsonify({"error": "Not found"}), 404
    data = request.get_json() or {}
    marks_map = data.get("marks") or {}
    feedback = (data.get("feedback") or "").strip()
    answers = db.execute(
        """SELECT a.*, q.marks AS max_marks FROM answers a
           JOIN questions q ON a.question_id = q.id WHERE a.submission_id = ?""",
        (sid,),
    ).fetchall()
    total = 0.0
    for a in answers:
        val = marks_map.get(str(a["id"]), marks_map.get(a["id"]))
        try:
            m = float(val) if val is not None and val != "" else (a["marks_awarded"] or 0)
            m = max(0, min(m, a["max_marks"]))
        except (TypeError, ValueError):
            m = a["marks_awarded"] or 0
        total += m
        db.execute("UPDATE answers SET marks_awarded = ? WHERE id = ?", (m, a["id"]))
    db.execute(
        """UPDATE submissions SET score=?, status='graded', graded_at=?, graded_by=?, feedback=?
           WHERE id=?""",
        (total, datetime.utcnow().isoformat(), current_user.id, feedback, sid),
    )
    db.commit()
    return jsonify({"message": "Graded", "score": total, "total_marks": sub["total_marks"]})


@app.route("/api/results/<int:sid>")
@login_required
def api_results(sid):
    db = get_db()
    sub = db.execute(
        """SELECT s.*, t.title AS test_title, t.total_marks, t.created_by, u.full_name AS student_name
           FROM submissions s
           JOIN tests t ON s.test_id = t.id
           JOIN users u ON s.student_id = u.id
           WHERE s.id = ?""",
        (sid,),
    ).fetchone()
    if not sub:
        return jsonify({"error": "Not found"}), 404
    if current_user.role == "student" and sub["student_id"] != current_user.id:
        return jsonify({"error": "Forbidden"}), 403
    if current_user.role == "teacher" and sub["created_by"] != current_user.id and sub["student_id"] != current_user.id:
        return jsonify({"error": "Forbidden"}), 403
    answers = db.execute(
        """SELECT a.*, q.question_text, q.question_type, q.marks AS max_marks, q.correct_answer
           FROM answers a JOIN questions q ON a.question_id = q.id
           WHERE a.submission_id = ?
           ORDER BY q.order_num, q.id""",
        (sid,),
    ).fetchall()
    return jsonify({"submission": row_to_dict(sub), "answers": [row_to_dict(a) for a in answers]})


# ── Admin users ────────────────────────────────────────────────────────────
@app.route("/api/users")
@login_required
@teacher_required
def api_users():
    db = get_db()
    rows = db.execute(
        "SELECT id, username, email, full_name, role, is_admin, created_at FROM users ORDER BY role, full_name"
    ).fetchall()
    return jsonify({"users": [row_to_dict(r) for r in rows]})


@app.route("/api/users", methods=["POST"])
@login_required
@admin_required
def api_create_teacher():
    """Only the default teacher/admin can create teacher accounts."""
    data = request.get_json() or {}
    username = (data.get("username") or "").strip()
    email = (data.get("email") or "").strip()
    full_name = (data.get("full_name") or "").strip()
    password = data.get("password") or ""
    if not all([username, email, full_name, password]):
        return jsonify({"error": "All fields are required"}), 400
    if len(password) < 6:
        return jsonify({"error": "Password must be at least 6 characters"}), 400
    db = get_db()
    try:
        db.execute(
            "INSERT INTO users (username, email, password_hash, full_name, role) VALUES (?,?,?,?,?)",
            (username, email, generate_password_hash(password), full_name, "teacher"),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": "Username or email already exists"}), 400
    return jsonify({"message": f"Teacher account created for {full_name}"})


def bootstrap():
    os.makedirs(os.path.join(app.config["UPLOAD_FOLDER"], "resources"), exist_ok=True)
    os.makedirs(os.path.dirname(app.config["DATABASE"]), exist_ok=True)
    with app.app_context():
        init_db()


# Runs on import so it also works under gunicorn on Render
bootstrap()


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=not IS_PRODUCTION, host="0.0.0.0", port=port)
